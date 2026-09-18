# mifi.py - ВАША ЛОГИКА + ФОРМАТИРОВАННЫЙ ЭКСПОРТ
import streamlit as st
import pandas as pd
import io
import re
from datetime import datetime, timedelta
import warnings
warnings.filterwarnings('ignore')

# ============================================================
# CSS СТИЛИ
# ============================================================

st.markdown("""
<style>
    .mifi-header {
        background: linear-gradient(135deg, #0a1628 0%, #1a2a4a 50%, #0d47a1 100%);
        padding: 24px 32px;
        border-radius: 14px;
        margin-bottom: 24px;
        box-shadow: 0 4px 20px rgba(13, 71, 161, 0.2);
    }
    .mifi-header h1 {
        color: #ffffff;
        font-family: 'Inter', sans-serif;
        font-size: 1.8rem;
        font-weight: 700;
        margin: 0;
        display: flex;
        align-items: center;
        gap: 12px;
    }
    .mifi-header p {
        color: rgba(255,255,255,0.7);
        font-family: 'Inter', sans-serif;
        margin: 4px 0 0 0;
        font-size: 0.95rem;
    }
    .mifi-header .badge-mifi {
        background: rgba(255,255,255,0.12);
        padding: 2px 14px;
        border-radius: 20px;
        font-size: 0.7rem;
        color: rgba(255,255,255,0.6);
        border: 1px solid rgba(255,255,255,0.08);
        margin-left: 12px;
    }
    .gradient-section {
        background: linear-gradient(135deg, #e3f2fd 0%, #f0f8ff 100%);
        padding: 16px 20px;
        border-radius: 12px;
        border-left: 4px solid #0d47a1;
        margin: 16px 0;
    }
    .gradient-section h3 {
        margin: 0;
        color: #0d47a1;
        font-family: 'Inter', sans-serif;
    }
    .gradient-section p {
        margin: 5px 0 0 0;
        color: #555;
        font-family: 'Inter', sans-serif;
    }
    .stButton button[kind="primary"] {
        background: linear-gradient(135deg, #0d47a1 0%, #1565c0 50%, #1a73e8 100%);
        color: #ffffff;
        border: none;
        box-shadow: 0 4px 16px rgba(13, 71, 161, 0.3);
    }
    .stButton button[kind="primary"]:hover {
        transform: translateY(-2px);
        box-shadow: 0 8px 32px rgba(13, 71, 161, 0.4);
    }
</style>
""", unsafe_allow_html=True)

# ============================================================
# ВАШИ ФУНКЦИИ (БЕЗ ИЗМЕНЕНИЙ)
# ============================================================

def parse_time_firebird(value):
    """Преобразование времени из Firebird (доля дня) в ЧЧ:ММ"""
    try:
        if pd.isna(value) or value == '' or value is None:
            return ''
        
        if isinstance(value, str) and re.match(r'^\d{2}:\d{2}$', value):
            return value
        
        if isinstance(value, (int, float)):
            if value > 0:
                total_seconds = int(round(value * 86400))
                hours = total_seconds // 3600
                minutes = (total_seconds % 3600) // 60
                return f"{hours:02d}:{minutes:02d}"
            return ''
        
        if isinstance(value, str):
            value = value.replace(',', '.')
            
            match = re.search(r'(\d{1,2}):(\d{2}):(\d{2})', value)
            if match:
                h = int(match.group(1))
                m = int(match.group(2))
                return f"{h:02d}:{m:02d}"
            
            match = re.search(r'(\d{1,2}):(\d{2})', value)
            if match:
                h = int(match.group(1))
                m = int(match.group(2))
                return f"{h:02d}:{m:02d}"
            
            match = re.search(r'^(\d+\.?\d*)$', value)
            if match:
                val = float(match.group(1))
                if val > 0:
                    total_seconds = int(round(val * 86400))
                    hours = total_seconds // 3600
                    minutes = (total_seconds % 3600) // 60
                    return f"{hours:02d}:{minutes:02d}"
            
            return value.strip()
        
        if hasattr(value, 'strftime'):
            return value.strftime('%H:%M')
        
        return str(value)
    except Exception as e:
        return str(value)

def parse_date_firebird(value):
    """Преобразование даты из Firebird в строку ДД.ММ.ГГГГ"""
    try:
        if pd.isna(value) or value == '' or value is None:
            return ''
        if isinstance(value, datetime):
            return value.strftime('%d.%m.%Y')
        if isinstance(value, (int, float)):
            try:
                dt = pd.to_datetime(value, unit='d', origin='1899-12-30')
                return dt.strftime('%d.%m.%Y')
            except:
                pass
        if isinstance(value, str):
            if re.match(r'^\d{2}\.\d{2}\.\d{4}$', value):
                return value
            match = re.search(r'(\d{2})\.(\d{2})\.(\d{4})', value)
            if match:
                return match.group(0)
            match = re.search(r'(\d{4})-(\d{2})-(\d{2})', value)
            if match:
                return f"{match.group(3)}.{match.group(2)}.{match.group(1)}"
            try:
                dt = pd.to_datetime(value, dayfirst=True)
                return dt.strftime('%d.%m.%Y')
            except:
                pass
        return str(value)
    except:
        return str(value)

def time_to_minutes(value):
    """Преобразование времени в минуты (поддерживает долю дня и строки)"""
    try:
        if pd.isna(value) or value == '' or value is None:
            return 0
        
        if isinstance(value, (int, float)):
            if value > 0:
                return int(round(value * 1440))
            return 0
        
        if isinstance(value, str):
            value = value.replace(',', '.')
            
            if ':' in value:
                parts = value.split(':')
                if len(parts) == 2:
                    return int(parts[0]) * 60 + int(parts[1])
                if len(parts) == 3:
                    return int(parts[0]) * 60 + int(parts[1])
            
            match = re.search(r'^(\d+\.?\d*)$', value)
            if match:
                val = float(match.group(1))
                if val < 1:
                    return int(round(val * 1440))
                else:
                    return int(round(val * 60))
        
        return 0
    except:
        return 0

def calculate_time_diff(time_in, time_out):
    """
    Расчет разницы между временем выхода и временем входа
    Возвращает строку в формате ЧЧ:ММ
    """
    try:
        min_in = time_to_minutes(time_in)
        min_out = time_to_minutes(time_out)
        
        if min_in == 0 and min_out == 0:
            return "00:00"
        
        diff_minutes = min_out - min_in
        if diff_minutes < 0:
            diff_minutes += 1440
        
        hours = diff_minutes // 60
        minutes = diff_minutes % 60
        return f"{hours:02d}:{minutes:02d}"
    except:
        return "00:00"

def calculate_time_diff_float(time_in, time_out):
    """
    Расчет разницы между временем выхода и временем входа
    Возвращает число с плавающей точкой (часы)
    """
    try:
        min_in = time_to_minutes(time_in)
        min_out = time_to_minutes(time_out)
        
        if min_in == 0 and min_out == 0:
            return 0.0
        
        diff_minutes = min_out - min_in
        if diff_minutes < 0:
            diff_minutes += 1440
        
        return round(diff_minutes / 60, 2)
    except:
        return 0.0

def parse_hours_firebird(value):
    """Преобразование часов в число с плавающей точкой"""
    try:
        if pd.isna(value) or value == '' or value is None:
            return 0.0
        
        if isinstance(value, (int, float)):
            return float(value)
        
        if isinstance(value, str):
            value = value.replace(',', '.')
            
            match = re.search(r'(\d{1,2}):(\d{2})', value)
            if match:
                h = int(match.group(1))
                m = int(match.group(2))
                return h + m / 60.0
            
            try:
                return float(value)
            except:
                pass
        
        return 0.0
    except:
        return 0.0

def format_hours_to_time(hours):
    """Преобразование часов (число) в формат ЧЧ:ММ"""
    if hours is None or hours == 0:
        return "00:00"
    
    hours = round(hours, 2)
    h = int(hours)
    minutes = int(round((hours - h) * 60))
    
    if minutes >= 60:
        h += 1
        minutes = 0
    
    return f"{h:02d}:{minutes:02d}"

# ============================================================
# ФУНКЦИЯ ЭКСПОРТА С ФОРМАТИРОВАНИЕМ (НОВАЯ)
# ============================================================

def export_to_excel_formatted(daily_df, summary_df):
    """
    Экспорт в Excel с форматированием:
    - Автоширина колонок
    - Заголовки: синий фон, белый жирный шрифт
    - Границы всех ячеек
    """
    output = io.BytesIO()
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        if daily_df is not None and not daily_df.empty:
            daily_export = daily_df.copy()
            
            if 'Время входа' in daily_export.columns:
                daily_export['Время входа'] = daily_export['Время входа'].apply(
                    lambda x: parse_time_firebird(x) if pd.notna(x) and x != '' else ''
                )
            if 'Время выхода' in daily_export.columns:
                daily_export['Время выхода'] = daily_export['Время выхода'].apply(
                    lambda x: parse_time_firebird(x) if pd.notna(x) and x != '' else ''
                )
            
            export_cols = ['Фамилия', 'Имя', 'Отчество', 'Дата', 'Время входа', 'Время выхода', 'Часов за день']
            daily_export[export_cols].to_excel(writer, sheet_name='По дню', index=False)
        
        summary_export = summary_df[['Фамилия', 'Имя', 'Отчество', 'Всего Ч']].copy()
        summary_export.to_excel(writer, sheet_name='Месяц всего', index=False)
    
    # Открываем workbook для форматирования
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    
    wb = load_workbook(output)
    
    # Стили
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    center_alignment = Alignment(horizontal='center', vertical='center')
    
    # Форматируем каждый лист
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Применяем стили к заголовкам (первая строка)
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = center_alignment
        
        # Применяем границы ко всем ячейкам с данными
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                if row > 1:  # Не перезаписываем стиль заголовка
                    cell.border = border
                    cell.alignment = center_alignment
        
        # Автоширина колонок
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
    
    wb.save(output)
    output.seek(0)
    return output.getvalue()

# ============================================================
# ФУНКЦИИ ПАРСИНГА (ВАШИ)
# ============================================================

def parse_firebird_file(uploaded_file):
    """Парсер для файлов Firebird"""
    debug_info = []
    
    try:
        df = pd.read_excel(uploaded_file, header=None)
        df = df.dropna(how='all')
        
        debug_info.append(f"📄 Файл прочитан, строк: {len(df)}")
        
        if df.empty:
            return None, None, "Файл пуст или не содержит данных", debug_info
        
        blocks = []
        current_block = None
        
        for idx, row in df.iterrows():
            if len(row) == 0 or pd.isna(row.iloc[0]):
                continue
            
            val = str(row.iloc[0]).strip()
            
            if val.isdigit() and len(val) <= 2:
                if current_block is not None and current_block['type'] != 'separator':
                    blocks.append(current_block)
                current_block = {'type': 'separator', 'start': idx, 'rows': []}
                continue
            
            if 'STAFF_ID' in val:
                if current_block is not None and current_block['type'] != 'separator':
                    blocks.append(current_block)
                current_block = {'type': 'staff', 'start': idx, 'rows': []}
                continue
            
            if 'FIRST_NAME' in val:
                if current_block is not None and current_block['type'] != 'separator':
                    blocks.append(current_block)
                current_block = {'type': 'first', 'start': idx, 'rows': []}
                continue
            
            if 'MIDDLE_NAME' in val:
                if current_block is not None and current_block['type'] != 'separator':
                    blocks.append(current_block)
                current_block = {'type': 'middle', 'start': idx, 'rows': []}
                continue
            
            if current_block is not None and current_block['type'] in ['staff', 'first', 'middle']:
                row_data = []
                for col in range(len(row)):
                    if pd.notna(row.iloc[col]):
                        row_data.append(str(row.iloc[col]).strip())
                    else:
                        row_data.append('')
                current_block['rows'].append(row_data)
        
        if current_block is not None:
            blocks.append(current_block)
        
        debug_info.append(f"🔍 Найдено блоков: {len(blocks)}")
        
        staff_blocks = [b for b in blocks if b['type'] == 'staff']
        first_blocks = [b for b in blocks if b['type'] == 'first']
        middle_blocks = [b for b in blocks if b['type'] == 'middle']
        
        debug_info.append(f"   - STAFF_ID: {len(staff_blocks)}")
        debug_info.append(f"   - FIRST_NAME: {len(first_blocks)}")
        debug_info.append(f"   - MIDDLE_NAME: {len(middle_blocks)}")
        
        employees = {}
        
        for mid_idx, mid_block in enumerate(middle_blocks):
            staff_block = staff_blocks[mid_idx] if mid_idx < len(staff_blocks) else None
            first_block = first_blocks[mid_idx] if mid_idx < len(first_blocks) else None
            
            last_names = []
            if staff_block is not None:
                for row in staff_block['rows']:
                    if len(row) > 0 and row[0]:
                        clean_val = row[0].replace('\xa0', '').replace(' ', '').replace('\u2009', '')
                        if clean_val.isdigit():
                            last_name = row[2] if len(row) > 2 and row[2] else ''
                            if last_name and last_name != 'nan':
                                last_names.append(last_name)
                            else:
                                last_name = row[1] if len(row) > 1 and row[1] else ''
                                if last_name and last_name != 'nan':
                                    last_names.append(last_name)
            
            first_names = []
            if first_block is not None:
                for row in first_block['rows']:
                    if len(row) > 0 and row[0] and not row[0].isdigit():
                        first_names.append(row[0])
            
            middle_data = []
            if mid_block is not None:
                for row in mid_block['rows']:
                    if len(row) > 0 and row[0] and not row[0].isdigit():
                        middle_name = row[0]
                        
                        date_val = row[4] if len(row) > 4 else ''
                        in_val = row[6] if len(row) > 6 else None
                        out_val = row[7] if len(row) > 7 else None
                        hours = parse_hours_firebird(row[8]) if len(row) > 8 else 0.0
                        
                        if hours == 0.0 and len(row) > 3 and row[3]:
                            hours = parse_hours_firebird(row[3])
                        
                        if date_val or hours > 0:
                            middle_data.append({
                                'middle': middle_name,
                                'date': date_val,
                                'in': in_val,
                                'out': out_val,
                                'hours': hours
                            })
            
            max_len = max(len(last_names), len(first_names), len(middle_data))
            
            for idx in range(max_len):
                last_name = last_names[idx] if idx < len(last_names) else ''
                first_name = first_names[idx] if idx < len(first_names) else ''
                
                if idx < len(middle_data):
                    item = middle_data[idx]
                    middle_name = item['middle']
                    
                    if last_name or first_name or middle_name:
                        key = (last_name, first_name, middle_name)
                        if key not in employees:
                            employees[key] = []
                        employees[key].append(item)
        
        debug_info.append(f"\n📊 Найдено сотрудников: {len(employees)}")
        
        result_records = []
        
        for (last_name, first_name, middle_name), records in employees.items():
            for record in records:
                date_val = record.get('date')
                if date_val:
                    date_formatted = parse_date_firebird(date_val)
                    if date_formatted:
                        time_in = record.get('in')
                        time_out = record.get('out')
                        
                        time_in_str = parse_time_firebird(time_in) if time_in is not None else ''
                        time_out_str = parse_time_firebird(time_out) if time_out is not None else ''
                        
                        time_diff_str = calculate_time_diff(time_in, time_out)
                        time_diff_float = calculate_time_diff_float(time_in, time_out)
                        
                        result_records.append({
                            'Фамилия': last_name,
                            'Имя': first_name,
                            'Отчество': middle_name,
                            'Дата': date_formatted,
                            'Время входа': time_in_str,
                            'Время выхода': time_out_str,
                            'Часов за день': time_diff_str,
                            'Часов за день (float)': time_diff_float
                        })
        
        debug_info.append(f"\n📋 Сформировано записей: {len(result_records)}")
        
        if not result_records:
            debug_info.append("❌ Не удалось сформировать ни одной записи!")
            return None, None, "Не удалось сформировать данные", debug_info
        
        result_df = pd.DataFrame(result_records)
        
        daily_df = result_df.drop_duplicates(subset=['Дата', 'Фамилия', 'Имя', 'Отчество', 'Время входа'])
        daily_df = daily_df[daily_df['Дата'].notna() & (daily_df['Дата'] != '')]
        
        if daily_df.empty:
            debug_info.append("❌ После фильтрации не осталось записей!")
            return None, None, "Не удалось сформировать данные", debug_info
        
        summary_df = daily_df.groupby(['Фамилия', 'Имя', 'Отчество']).agg({
            'Часов за день (float)': 'sum'
        }).reset_index()
        summary_df.columns = ['Фамилия', 'Имя', 'Отчество', 'Всего часов (float)']
        summary_df['Всего Ч'] = summary_df['Всего часов (float)'].apply(format_hours_to_time)
        summary_df = summary_df.sort_values('Всего часов (float)', ascending=False).reset_index(drop=True)
        
        debug_info.append(f"✅ Детальный отчет: {len(daily_df)} записей, {len(summary_df)} сотрудников")
        return daily_df, summary_df, None, debug_info
        
    except Exception as e:
        debug_info.append(f"❌ Ошибка: {str(e)}")
        import traceback
        debug_info.append(traceback.format_exc())
        return None, None, f"Ошибка при обработке: {str(e)}", debug_info

def merge_two_files(file1, file2):
    """Объединение двух файлов табеля"""
    try:
        daily1, summary1, error1, debug1 = parse_firebird_file(file1)
        if error1:
            return None, None, f"Ошибка в первом файле: {error1}", debug1
        
        daily2, summary2, error2, debug2 = parse_firebird_file(file2)
        if error2:
            return None, None, f"Ошибка во втором файле: {error2}", debug2
        
        if daily1 is not None and daily2 is not None:
            daily_merged = pd.concat([daily1, daily2], ignore_index=True)
            daily_merged = daily_merged.drop_duplicates(subset=['Дата', 'Фамилия', 'Имя', 'Отчество', 'Время входа'])
            daily_merged = daily_merged.sort_values(['Дата', 'Фамилия', 'Имя', 'Отчество']).reset_index(drop=True)
            
            summary_merged = daily_merged.groupby(['Фамилия', 'Имя', 'Отчество']).agg({
                'Часов за день (float)': 'sum'
            }).reset_index()
            summary_merged.columns = ['Фамилия', 'Имя', 'Отчество', 'Всего часов (float)']
            summary_merged['Всего Ч'] = summary_merged['Всего часов (float)'].apply(format_hours_to_time)
            summary_merged = summary_merged.sort_values('Всего часов (float)', ascending=False).reset_index(drop=True)
            
            return daily_merged, summary_merged, None, debug1 + debug2
        else:
            return None, None, "Файлы имеют несовместимый формат", debug1 + debug2
        
    except Exception as e:
        return None, None, f"Ошибка при объединении: {str(e)}", []

# ============================================================
# ОСНОВНАЯ ФУНКЦИЯ МОДУЛЯ
# ============================================================

def main():
    """Главная функция модуля МИФИ"""
    
    st.markdown("""
    <div class="mifi-header">
        <h1>
            🏛️ МИФИ
            <span class="badge-mifi">Технологический институт</span>
        </h1>
        <p>ТИ НИЯУ МИФИ — табель учета рабочего времени</p>
    </div>
    """, unsafe_allow_html=True)
    
    st.markdown("""
    <div class="gradient-section">
        <h3>🔄 Преобразование табеля из Firebird</h3>
        <p>Загрузите файл(ы) табеля из Firebird и получите преобразованные данные</p>
    </div>
    """, unsafe_allow_html=True)
    
    st.subheader("📂 Загрузка файлов")
    
    col1, col2 = st.columns(2)
    
    file1 = None
    file2 = None
    
    with col1:
        st.markdown("**Файл 1**")
        file1 = st.file_uploader(
            "Выберите первый файл",
            type=['xls', 'xlsx'],
            key="mifi_uploader_1",
            label_visibility="collapsed"
        )
        if file1:
            st.success(f"✅ Загружен: {file1.name}")
    
    with col2:
        st.markdown("**Файл 2 (необязательно)**")
        file2 = st.file_uploader(
            "Выберите второй файл",
            type=['xls', 'xlsx'],
            key="mifi_uploader_2",
            label_visibility="collapsed"
        )
        if file2:
            st.success(f"✅ Загружен: {file2.name}")
    
    if file1 is not None:
        if st.button("🔄 Обработать файлы", use_container_width=True, type="primary"):
            with st.spinner("⏳ Обработка файлов..."):
                all_debug = []
                
                if file2 is not None:
                    daily_df, summary_df, error, debug = merge_two_files(file1, file2)
                    all_debug = debug
                    if error:
                        st.error(f"❌ {error}")
                    else:
                        st.success("✅ Файлы успешно объединены и обработаны!")
                else:
                    daily_df, summary_df, error, debug = parse_firebird_file(file1)
                    all_debug = debug
                    if error:
                        st.error(f"❌ {error}")
                    else:
                        st.success("✅ Файл успешно обработан!")
                
                with st.expander("🔍 Отладочная информация"):
                    for line in all_debug:
                        st.text(line)
                
                if summary_df is not None and not summary_df.empty:
                    st.session_state.mifi_summary = summary_df
                    
                    if daily_df is not None and not daily_df.empty:
                        st.session_state.mifi_daily = daily_df
                        col1, col2, col3 = st.columns(3)
                        with col1:
                            st.metric("📋 Всего записей", len(daily_df))
                        with col2:
                            st.metric("👥 Сотрудников", len(summary_df))
                        with col3:
                            total_hours = summary_df['Всего часов (float)'].sum()
                            st.metric("⏱ Всего часов", format_hours_to_time(total_hours))
                        
                        st.subheader("📋 Детальные записи (первые 10 строк)")
                        display_cols = ['Фамилия', 'Имя', 'Отчество', 'Дата', 'Время входа', 'Время выхода', 'Часов за день']
                        st.dataframe(daily_df[display_cols].head(10), use_container_width=True, hide_index=True)
                    else:
                        col1, col2 = st.columns(2)
                        with col1:
                            st.metric("👥 Сотрудников", len(summary_df))
                        with col2:
                            total_hours = summary_df['Всего часов (float)'].sum()
                            st.metric("⏱ Всего часов", format_hours_to_time(total_hours))
                        
                        st.info("📊 Это сводный файл (SUM_HOURS). Отображаются только итоговые часы за месяц.")
                    
                    st.subheader("📊 Сводка по сотрудникам")
                    display_cols = ['Фамилия', 'Имя', 'Отчество', 'Всего Ч']
                    st.dataframe(summary_df[display_cols], use_container_width=True, hide_index=True)
                    
                    st.subheader("📥 Экспорт результатов")
                    
                    col1, col2 = st.columns(2)
                    
                    with col1:
                        # Используем новую функцию с форматированием
                        excel_data = export_to_excel_formatted(daily_df, summary_df)
                        
                        st.download_button(
                            label="📥 Скачать преобразованный Excel",
                            data=excel_data,
                            file_name=f"Табель_преобразованный_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True
                        )
                    
                    with col2:
                        csv_export = summary_df[['Фамилия', 'Имя', 'Отчество', 'Всего Ч']].copy()
                        st.download_button(
                            label="📥 Скачать сводку CSV",
                            data=csv_export.to_csv(index=False, encoding='utf-8-sig'),
                            file_name=f"Сводка_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
                            mime="text/csv",
                            use_container_width=True
                        )
                    
                    st.info("💡 Результат можно использовать для дальнейшего анализа и построения отчетов.")
                else:
                    st.warning("⚠️ Не удалось получить данные для отображения")
    
    st.markdown("---")
    st.info("ℹ️ Здесь может быть остальная функциональность модуля МИФИ")

if __name__ == "__main__":
    main()